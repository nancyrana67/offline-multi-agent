"""Document processor for RAG system."""
import logging
import os
from pathlib import Path
from typing import List, Dict, Any, Optional
import PyPDF2
from docx import Document

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Processes various document formats."""

    SUPPORTED_FORMATS = {".pdf", ".docx", ".txt", ".md"}

    @staticmethod
    def extract_text(file_path: str) -> Optional[str]:
        """Extract text from various file formats."""
        try:
            file_ext = Path(file_path).suffix.lower()

            if file_ext == ".pdf":
                return DocumentProcessor._extract_pdf(file_path)
            elif file_ext == ".docx":
                return DocumentProcessor._extract_docx(file_path)
            elif file_ext == ".txt":
                return DocumentProcessor._extract_text(file_path)
            elif file_ext == ".md":
                return DocumentProcessor._extract_text(file_path)
            else:
                logger.warning(f"Unsupported file format: {file_ext}")
                return None

        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {str(e)}")
            return None

    @staticmethod
    def _extract_pdf(file_path: str) -> str:
        """Extract text from PDF."""
        text = ""
        try:
            with open(file_path, "rb") as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text()
        except Exception as e:
            logger.error(f"PDF extraction error: {str(e)}")
        return text

    @staticmethod
    def _extract_docx(file_path: str) -> str:
        """Extract text from DOCX."""
        text = ""
        try:
            doc = Document(file_path)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
        except Exception as e:
            logger.error(f"DOCX extraction error: {str(e)}")
        return text

    @staticmethod
    def _extract_text(file_path: str) -> str:
        """Extract text from TXT or MD."""
        try:
            with open(file_path, "r", encoding="utf-8") as file:
                return file.read()
        except Exception as e:
            logger.error(f"Text extraction error: {str(e)}")
            return ""

    @staticmethod
    def chunk_text(
        text: str, chunk_size: int = 500, overlap: int = 50
    ) -> List[str]:
        """Split text into overlapping chunks."""
        chunks = []
        text_length = len(text)

        for i in range(0, text_length, chunk_size - overlap):
            chunk = text[i : i + chunk_size]
            if len(chunk.strip()) > 0:
                chunks.append(chunk)

            if i + chunk_size >= text_length:
                break

        return chunks

    @staticmethod
    def clean_text(text: str) -> str:
        """Clean and normalize text."""
        # Remove extra whitespace
        text = " ".join(text.split())
        # Remove special characters (keep alphanumeric, spaces, and basic punctuation)
        import re
        text = re.sub(r"[^\w\s.!?,-]", "", text)
        return text.strip()

    @staticmethod
    def get_document_metadata(file_path: str) -> Dict[str, Any]:
        """Extract metadata from document."""
        try:
            file_path = Path(file_path)
            return {
                "filename": file_path.name,
                "path": str(file_path),
                "size": file_path.stat().st_size,
                "created": file_path.stat().st_ctime,
                "modified": file_path.stat().st_mtime,
                "format": file_path.suffix.lower(),
            }
        except Exception as e:
            logger.error(f"Error getting metadata: {str(e)}")
            return {}
